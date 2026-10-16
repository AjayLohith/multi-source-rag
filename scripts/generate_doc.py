import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def markdown_to_docx(md_path, docx_path):
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal in_table, table_lines
        if not table_lines:
            in_table = False
            return
        
        # Parse table rows
        rows_data = []
        for tl in table_lines:
            tl = tl.strip()
            if not tl or re.match(r"^\|?\s*[-:]+[-| :]*\s*\|?$", tl):
                continue
            cells = [c.strip() for c in tl.strip("|").split("|")]
            rows_data.append(cells)

        if rows_data:
            num_cols = max(len(r) for r in rows_data)
            # Normalize column counts
            for r in rows_data:
                while len(r) < num_cols:
                    r.append("")

            table = doc.add_table(rows=len(rows_data), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True

            for r_idx, row in enumerate(rows_data):
                is_header = (r_idx == 0)
                for c_idx, cell_value in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    run = p.add_run(cell_value.replace("**", ""))
                    if is_header:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        set_cell_background(cell, "1F4E79") # Deep navy
                    else:
                        if "**" in cell_value:
                            run.bold = True
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F2F4F7") # Light zebra stripe
                        else:
                            set_cell_background(cell, "FFFFFF")
                    set_cell_margins(cell, top=120, bottom=120, left=140, right=140)

            # Add border formatting to table
            tblPr = table._tbl.tblPr
            borders = parse_xml(
                f'<w:tblBorders {nsdecls("w")}>'
                f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
                f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="1F4E79"/>'
                f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
                f'<w:insideV w:val="none"/>'
                f'</w:tblBorders>'
            )
            tblPr.append(borders)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)

        table_lines = []
        in_table = False

    def flush_code():
        nonlocal in_code_block, code_lines
        if not code_lines:
            in_code_block = False
            return
        code_text = "".join(code_lines).rstrip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.right_indent = Inches(0.2)
        
        # Add single cell table with shaded background for code
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8F9FA")
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        
        # Code borders
        tblPr = tbl._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'<w:left w:val="single" w:sz="12" w:space="0" w:color="3182CE"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(0)
        cp.paragraph_format.space_after = Pt(0)
        crun = cp.add_run(code_text)
        crun.font.name = 'Consolas'
        crun.font.size = Pt(9.5)
        crun.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle code blocks
        if stripped.startswith("```"):
            if in_table:
                flush_table()
            if in_code_block:
                flush_code()
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            table_lines.append(stripped)
            i += 1
            continue
        elif in_table:
            flush_table()

        # Handle horizontal divider
        if stripped in ["---", "***", "___"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(title_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79) # Executive navy
            i += 1
            continue

        if stripped.startswith("## "):
            h2_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h2_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0) # Steel blue
            i += 1
            continue

        if stripped.startswith("### "):
            h3_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h3_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68) # Charcoal
            i += 1
            continue

        # Bullet list
        if stripped.startswith("- ") or stripped.startswith("* "):
            bullet_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            _format_runs(p, bullet_text)
            i += 1
            continue

        # Numbered list
        m_num = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_num:
            num_text = m_num.group(2).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            _format_runs(p, num_text)
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        _format_runs(p, stripped)
        i += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    doc.save(docx_path)
    print(f"Generated successfully: {docx_path}")

def _format_runs(paragraph, text):
    # Parses bold **text**, code `text`, and plain text
    tokens = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x80, 0x20, 0x20)
        else:
            paragraph.add_run(token)

if __name__ == "__main__":
    src = r"d:\multi-source-rag-app\TECHNICAL_DOCS.md"
    docx_dest = r"d:\multi-source-rag-app\TECHNICAL_DOCS.docx"
    doc_dest = r"d:\multi-source-rag-app\TECHNICAL_DOCS.doc"
    markdown_to_docx(src, docx_dest)
    
    # Also create the .doc copy for software or preferences expecting .doc extension
    import shutil
    shutil.copyfile(docx_dest, doc_dest)
    print(f"Copied to: {doc_dest}")
