from pathlib import Path
from typing import Any, Dict
import docx

from app.extraction.errors import DocExtractionError


def extract_docx(file_path: str) -> Dict[str, Any]:
    """
    Extracts text content and paragraph/table elements from Word (.docx) files.
    """
    path = Path(file_path)
    if not path.exists():
        raise DocExtractionError(
            message=f"Docx file not found: {file_path}",
            source_ref=str(file_path),
            error_code="DOCX_NOT_FOUND",
        )

    try:
        doc = docx.Document(str(path))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_texts.append(row_str)

        all_text = "\n\n".join(paragraphs + table_texts)
        return {
            "file_path": str(path),
            "text": all_text,
            "paragraph_count": len(paragraphs),
            "table_count": len(doc.tables),
        }
    except Exception as exc:
        raise DocExtractionError(
            message=f"Failed to extract docx document: {exc}",
            source_ref=str(file_path),
            error_code="DOCX_PARSE_ERROR",
        ) from exc
